"""Local document extraction with page-level OCR and explicit failure modes."""
import os, shutil, subprocess, tempfile
from pathlib import Path
from io import BytesIO
from pypdf import PdfReader
from docx import Document


def run(args, timeout=120):
    return subprocess.run(args, check=True, capture_output=True, timeout=timeout).stdout


def extract_pages(name, content):
    ext = Path(name).suffix.lower()
    if ext == '.txt':
        return [{'page': 1, 'text': content.decode('utf-8', errors='replace'), 'method': 'text'}]
    if ext == '.docx':
        doc = Document(BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        parts += [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        return [{'page': 1, 'text': '\n'.join(parts), 'method': 'docx', 'pagination': 'logical_not_physical'}]
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / ('source' + ext)
        source.write_bytes(content)
        if ext == '.doc':
            if not shutil.which('antiword'):
                raise ValueError('Lecture .doc indisponible : installer antiword ou convertir en DOCX.')
            text = run(['antiword', '-w', '0', str(source)]).decode('utf-8', errors='replace')
            return [{'page': i+1, 'text': part, 'method': 'antiword', 'pagination': 'logical_not_physical'} for i, part in enumerate(text.split('\f'))]
        if ext != '.pdf':
            raise ValueError('Format non pris en charge')
        reader = PdfReader(BytesIO(content))
        limit = int(os.getenv('MAX_PAGES', '250'))
        if len(reader.pages) > limit:
            raise ValueError(f'Document trop long : limite {limit} pages.')
        pages = []
        langs = ''
        for index, page in enumerate(reader.pages):
            text = (page.extract_text() or '').strip()
            method, warning = 'pdf_text', ''
            if len(text) < 60:
                if not all(shutil.which(tool) for tool in ('pdftoppm', 'tesseract')):
                    warning = 'OCR requis : installer poppler-utils et tesseract-ocr.'
                else:
                    if not langs:
                        installed = run(['tesseract', '--list-langs']).decode()
                        langs = 'fra+eng' if 'fra' in installed else 'eng'
                    prefix = str(Path(directory) / f'page_{index+1}')
                    run(['pdftoppm', '-f', str(index+1), '-l', str(index+1), '-scale-to', '2400', '-singlefile', '-png', str(source), prefix])
                    text = run(['tesseract', prefix+'.png', 'stdout', '-l', langs]).decode('utf-8', errors='replace').strip()
                    if len(text) < 300:
                        alternative = run(['tesseract', prefix+'.png', 'stdout', '-l', langs, '--psm', '11']).decode('utf-8', errors='replace').strip()
                        if len(alternative) > len(text): text = alternative
                    method = 'ocr_'+langs
                    warning = 'Lecture OCR à contrôler visuellement : noms, dates et diplômes.'
            if not text:
                warning += ' AUCUN TEXTE LISIBLE : contrôle manuel obligatoire.'
            pages.append({'page': index+1, 'text': text, 'method': method, 'warning': warning})
        return pages
