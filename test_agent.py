import os, tempfile, unittest
os.environ['DATA_DIR']=tempfile.mkdtemp()
from fastapi.testclient import TestClient
from app import app, verify_quotes, make_draft
from document_reader import extract_pages
from docx import Document

class AgentTests(unittest.TestCase):
    def test_health(self):
        self.assertEqual(TestClient(app).get('/health').status_code,200)
    def test_invalid_id(self):
        self.assertEqual(TestClient(app).post('/api/cases/bad/confirm',json={'candidate_confirmed':True}).status_code,400)
    def test_unverified_match_downgraded(self):
        result={'dao_requirements':[{'source_id':'S01','page':1,'quote':'absent'}], 'matches':[{'status':'conforme'}]}
        self.assertEqual(verify_quotes(result,{'S01':[{'page':1,'text':'BTS'}]})['matches'][0]['status'],'non_évaluable')
    def test_no_invented_experience_export(self):
        case={'candidate':{'candidate_name':'TEST'},'analysis':{'candidate_facts':[{'field':'expérience','value':'CHANTIER_NON_CONFIRME','status':'à_confirmer'}]}}
        doc=Document(make_draft(case))
        self.assertNotIn('CHANTIER_NON_CONFIRME',' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells))
    def test_without_provider(self):
        os.environ['MODEL_PROVIDER']=''
        response=TestClient(app).post('/api/analyze',data={'candidate_name':'TEST'},files={'dao':('dao.txt',b'Personnel BTS Batiment')})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['status'],'needs_model')
    def test_team_invalid_mapping(self):
        response=TestClient(app).post('/api/team/analyze',data={'roster':'[{"name":"A","documents":["missing.pdf"]}]'},files=[('dao',('dao.txt',b'DAO')),('documents',('a.txt',b'A'))])
        self.assertEqual(response.status_code,400)
    def test_text(self):
        self.assertEqual(extract_pages('a.txt',b'BTS')[0]['text'],'BTS')
if __name__=='__main__': unittest.main()

class PeriodTests(unittest.TestCase):
    def test_overlap_and_internship(self):
        from compliance import relevant_months
        periods=[{'from':'2020-01-01','to':'2021-01-01','status':'prouvé'}, {'from':'2020-06-01','to':'2021-01-01','status':'prouvé'}, {'from':'2019-01-01','to':'2020-01-01','status':'prouvé','internship':True}]
        self.assertEqual(relevant_months(periods,'2022-01-01'),12)
    def test_future_cap(self):
        from compliance import relevant_months
        self.assertEqual(relevant_months([{'from':'2020-01-01','to':'2025-01-01','status':'prouvé'}],'2021-01-01'),12)

class ExportTests(unittest.TestCase):
    def test_company_separated(self):
        case={'candidate':{'candidate_name':'PERSONNE','company':'ENTREPRISE'},'analysis':{}}
        text=' '.join(p.text for p in Document(make_draft(case)).paragraphs)
        self.assertIn('Nom du candidat (entreprise) : ENTREPRISE',text)
        self.assertIn('Nom et prénoms : PERSONNE',text)
    def test_team_export_requires_confirmation(self):
        client=TestClient(app)
        os.environ['MODEL_PROVIDER']=''
        case=client.post('/api/analyze',data={'candidate_name':'PERSONNE','company':'ENTREPRISE'},files={'dao':('dao.txt',b'DAO')}).json()
        response=client.post('/api/team/per1',json={'company':'ENTREPRISE','case_ids':[case['case_id']]})
        self.assertEqual(response.status_code,409)
